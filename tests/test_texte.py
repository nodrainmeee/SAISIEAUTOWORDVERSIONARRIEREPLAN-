"""Tests du traitement du texte et de l'importation (.txt, .docx)."""

import os
import random
import tempfile
import unittest

from autosaisie import libelles as L
from autosaisie.rythme import MODE_HUMAIN, MODE_RAPIDE, delai_apres, prochain_groupe
from autosaisie.texte import (ErreurImportation, importer_fichier, lire_docx, lire_txt,
                              normaliser_fins_de_ligne, texte_est_vide)

TEXTE_FR = "Été à Cotonou : où êtes-vous ? « Çà et là », œuvre naïve.\nDeuxième ligne.\n\nTroisième."


class TestTexte(unittest.TestCase):
    def setUp(self):
        self.dossier = tempfile.TemporaryDirectory()
        self.addCleanup(self.dossier.cleanup)

    def chemin(self, nom):
        return os.path.join(self.dossier.name, nom)

    def ecrire(self, nom, octets):
        with open(self.chemin(nom), "wb") as f:
            f.write(octets)
        return self.chemin(nom)

    def test_txt_utf8_accents_et_retours_a_la_ligne(self):
        p = self.ecrire("a.txt", TEXTE_FR.encode("utf-8"))
        self.assertEqual(lire_txt(p), TEXTE_FR)

    def test_txt_utf8_avec_bom(self):
        p = self.ecrire("b.txt", b"\xef\xbb\xbf" + TEXTE_FR.encode("utf-8"))
        self.assertEqual(lire_txt(p), TEXTE_FR)

    def test_txt_windows_1252(self):
        p = self.ecrire("c.txt", "Café été".encode("cp1252"))
        self.assertEqual(lire_txt(p), "Café été")

    def test_txt_utf16(self):
        p = self.ecrire("d.txt", TEXTE_FR.encode("utf-16"))
        self.assertEqual(lire_txt(p), TEXTE_FR)

    def test_fins_de_ligne_windows_normalisees(self):
        p = self.ecrire("e.txt", b"un\r\ndeux\r\ntrois")
        self.assertEqual(lire_txt(p), "un\ndeux\ntrois")
        self.assertEqual(normaliser_fins_de_ligne("a\rb"), "a\nb")

    def test_fichier_inexistant_message_francais(self):
        with self.assertRaises(ErreurImportation) as ctx:
            lire_txt(self.chemin("absent.txt"))
        self.assertEqual(str(ctx.exception), L.MSG_FICHIER_ILLISIBLE)

    def test_extension_non_prise_en_charge(self):
        p = self.ecrire("f.pdf", b"x")
        with self.assertRaises(ErreurImportation):
            importer_fichier(p)

    def test_texte_vide(self):
        self.assertTrue(texte_est_vide(""))
        self.assertTrue(texte_est_vide("  \n\t "))
        self.assertFalse(texte_est_vide("a"))

    def test_docx_paragraphes_et_tableaux_dans_l_ordre(self):
        import docx
        d = docx.Document()
        d.add_paragraph("Avant le tableau : é à ç")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "A1"
        t.cell(0, 1).text = "B1"
        t.cell(1, 0).text = "A2"
        t.cell(1, 1).text = "B2"
        d.add_paragraph("Après le tableau")
        p = self.chemin("t.docx")
        d.save(p)
        self.assertEqual(lire_docx(p), "Avant le tableau : é à ç\nA1\tB1\nA2\tB2\nAprès le tableau")
        self.assertEqual(importer_fichier(p), lire_docx(p))

    def test_docx_corrompu(self):
        p = self.ecrire("mauvais.docx", b"ceci n'est pas un docx")
        with self.assertRaises(ErreurImportation) as ctx:
            lire_docx(p)
        self.assertEqual(str(ctx.exception), L.MSG_FICHIER_ILLISIBLE)


class TestRythme(unittest.TestCase):
    def test_groupes_reconstituent_le_texte_sans_perte(self):
        for taille in (1, 3, 4, 10):
            pos, sortie = 0, []
            while pos < len(TEXTE_FR):
                g = prochain_groupe(TEXTE_FR, pos, taille)
                self.assertTrue(0 < len(g) <= taille)
                sortie.append(g)
                pos += len(g)
            self.assertEqual("".join(sortie), TEXTE_FR)

    def test_retour_a_la_ligne_envoye_seul(self):
        self.assertEqual(prochain_groupe("ab\ncd", 2, 5), "\n")
        self.assertEqual(prochain_groupe("ab\ncd", 0, 5), "ab")

    def test_mode_rapide_delai_regulier(self):
        a = random.Random(1)
        self.assertAlmostEqual(delai_apres("abcd", MODE_RAPIDE, 40, 5, a), 0.1)

    def test_mode_humain_varie_et_ajoute_des_pauses(self):
        a = random.Random(2)
        delais = [delai_apres("a", MODE_HUMAIN, 10, 5, a) for _ in range(200)]
        self.assertGreater(max(delais), min(delais))
        self.assertGreater(delai_apres(".", MODE_HUMAIN, 10, 10, random.Random(3)), 0.1)

    def test_intensite_zero_pas_de_pause_ajoutee(self):
        a = random.Random(4)
        for _ in range(100):
            self.assertLessEqual(delai_apres(".", MODE_HUMAIN, 10, 0, a), 0.1 * 1.5 + 1e-9)


class TestLibelles(unittest.TestCase):
    def test_libelles_obligatoires_presents(self):
        attendus = {
            "BTN_IMPORT_TXT": "Importer un fichier texte",
            "BTN_IMPORT_DOCX": "Importer un document Word",
            "BTN_DEMARRER": "Démarrer la saisie",
            "BTN_PAUSE": "Mettre en pause",
            "BTN_REPRENDRE": "Reprendre la saisie",
            "BTN_ARRETER": "Arrêter la saisie",
            "BTN_EFFACER": "Effacer le texte",
            "BTN_IMPORTER": "Importer un fichier",
            "BTN_PARAMETRES": "Paramètres",
            "BTN_APROPOS": "À propos",
            "BTN_QUITTER": "Quitter",
            "MODE_RAPIDE": "Mode rapide",
            "MODE_HUMAIN": "Mode humain",
            "WORD_DETECTE": "Microsoft Word détecté",
            "WORD_NON_DETECTE": "Microsoft Word non détecté",
        }
        for nom, valeur in attendus.items():
            self.assertEqual(getattr(L, nom), valeur, nom)

    def test_aucun_mot_anglais_courant_dans_les_libelles(self):
        interdits = (" start", "stop ", "pause button", "error", "loading", "cancel", "settings", "import file")
        for nom in dir(L):
            valeur = getattr(L, nom)
            if nom.isupper() and isinstance(valeur, str):
                for mot in interdits:
                    self.assertNotIn(mot, " " + valeur.lower() + " ", nom)

    def test_messages_securite_exacts(self):
        self.assertTrue(L.MSG_WORD_REQUIS.startswith("Saisie impossible : Microsoft Word doit être la fenêtre active."))
        self.assertTrue(L.MSG_SUSPENDU.startswith("Saisie suspendue pour votre sécurité"))


if __name__ == "__main__":
    unittest.main()
